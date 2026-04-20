# -*- coding: utf-8 -*-
"""
Генератор Word-документа курсовой работы для проекта FitnessCenter.
Запуск: python generate_docx.py
Результат: e:\Fitness_web\docs\Курсовая_FitnessCenter.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

BASE_DIR = r"e:\Fitness_web"
DOCS_DIR = os.path.join(BASE_DIR, "docs")
IMAGES_DIR = os.path.join(DOCS_DIR, "images")
OUTPUT_PATH = os.path.join(DOCS_DIR, "Курсовая_FitnessCenter.docx")

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
HEADING1_SIZE = Pt(16)
HEADING3_SIZE = Pt(14)
LINE_SPACING = 1.5


# ─── Helpers ─────────────────────────────────────────────────────────────

def set_run_font(run, name=FONT_NAME, size=FONT_SIZE, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def add_paragraph(doc, text, style="Normal", alignment=None, bold=False,
                  italic=False, size=None, space_after=None, space_before=None,
                  first_line_indent=None, line_spacing=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_run_font(run, size=size or FONT_SIZE, bold=bold, italic=italic)
    if alignment is not None:
        p.alignment = alignment
    fmt = p.paragraph_format
    if space_after is not None:
        fmt.space_after = space_after
    if space_before is not None:
        fmt.space_before = space_before
    if first_line_indent is not None:
        fmt.first_line_indent = first_line_indent
    if line_spacing:
        fmt.line_spacing = line_spacing
    else:
        fmt.line_spacing = LINE_SPACING
    return p


def add_heading1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=HEADING1_SIZE, bold=True)
    fmt = p.paragraph_format
    fmt.space_before = Pt(12)
    fmt.space_after = Pt(6)
    fmt.line_spacing = LINE_SPACING
    fmt.keep_with_next = True
    # Set outline level for TOC
    p.style = doc.styles["Heading 1"]
    for r in p.runs:
        set_run_font(r, size=HEADING1_SIZE, bold=True)
    return p


def add_heading3(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 3"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.runs[0] if p.runs else p.add_run(text)
    if not p.runs or p.runs[0].text != text:
        p.clear()
        run = p.add_run(text)
    set_run_font(run, size=HEADING3_SIZE, bold=True)
    fmt = p.paragraph_format
    fmt.space_before = Pt(10)
    fmt.space_after = Pt(4)
    fmt.line_spacing = LINE_SPACING
    fmt.first_line_indent = Cm(1.25)
    return p


def add_body(doc, text, indent=True):
    return add_paragraph(
        doc, text,
        first_line_indent=Cm(1.25) if indent else None,
        space_after=Pt(0),
        space_before=Pt(0),
    )


def add_figure(doc, image_path, caption, fig_num, width=None):
    """Insert image centered with caption below."""
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.line_spacing = 1.0

    if os.path.exists(image_path):
        if width:
            p_img.add_run().add_picture(image_path, width=width)
        else:
            p_img.add_run().add_picture(image_path, width=Cm(15))
    else:
        run = p_img.add_run(f"[Изображение не найдено: {image_path}]")
        set_run_font(run, italic=True, color=RGBColor(255, 0, 0))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(6)
    p_cap.paragraph_format.line_spacing = LINE_SPACING
    run = p_cap.add_run(f"Рисунок {fig_num}. {caption}")
    set_run_font(run, size=Pt(12), italic=True)
    return fig_num + 1


def add_screenshot_placeholder(doc, caption, fig_num):
    """Add bordered placeholder for a screenshot."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = LINE_SPACING

    # Create bordered box via a 1x1 table
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cp.add_run("\n\n[Место для скриншота]\n\n")
    set_run_font(run, size=Pt(12), italic=True, color=RGBColor(150, 150, 150))

    # Set cell borders
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="999999"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="999999"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="999999"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="999999"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)
    # Set cell width
    cell.width = Cm(15)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(6)
    p_cap.paragraph_format.line_spacing = LINE_SPACING
    run = p_cap.add_run(f"Рисунок {fig_num}. {caption}")
    set_run_font(run, size=Pt(12), italic=True)
    return fig_num + 1


def add_table_from_data(doc, headers, rows):
    """Create a bordered table with header row shaded."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size=Pt(12), bold=True)
        p.paragraph_format.line_spacing = 1.0
        # Shade header
        shading = parse_xml(
            f'<w:shd {nsdecls("w")} w:fill="D9D9D9" w:val="clear"/>'
        )
        cell._tc.get_or_add_tcPr().append(shading)

    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=Pt(12))
            p.paragraph_format.line_spacing = 1.0

    # Add space after table
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    return table


def add_toc(doc):
    """Insert a Table of Contents field."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._r.append(fldChar1)

    run2 = p.add_run()
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
    run2._r.append(instrText)

    run3 = p.add_run()
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    run3._r.append(fldChar2)

    run4 = p.add_run("(Обновите оглавление: ПКМ → Обновить поле)")
    set_run_font(run4, size=Pt(12), italic=True, color=RGBColor(128, 128, 128))

    run5 = p.add_run()
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run5._r.append(fldChar3)


# ─── Document Construction ───────────────────────────────────────────────

def build_document():
    doc = Document()
    fig_num = 1

    # ── Page setup ──
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    # ── Default style ──
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE
    style.paragraph_format.line_spacing = LINE_SPACING
    rpr = style.element.rPr
    if rpr is None:
        style.element.get_or_add_rPr()
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # Configure Heading 1
    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_NAME
    h1.font.size = HEADING1_SIZE
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.line_spacing = LINE_SPACING
    h1.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # Configure Heading 3
    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_NAME
    h3.font.size = HEADING3_SIZE
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.line_spacing = LINE_SPACING
    h3.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # ════════════════════════════════════════════════════════════════════
    # TITLE PAGE
    # ════════════════════════════════════════════════════════════════════
    for _ in range(2):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_paragraph(doc, "УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "«УНИВЕРСИТЕТ УПРАВЛЕНИЯ «ТИСБИ»",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "Факультет информационных технологий",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "Кафедра информационных технологий",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE, space_after=Pt(0))

    for _ in range(4):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_paragraph(doc, "Курсовая работа",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(18), space_after=Pt(0))
    add_paragraph(doc, "по дисциплине «Разработка Интернет-приложений»",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "на тему: «Разработка веб-приложения фитнес-центра",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "с административной панелью и личным кабинетом»",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE, space_after=Pt(0))

    for _ in range(5):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_paragraph(doc, "Выполнила:",
                  alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "ст. группы П-111",
                  alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "Шипицын А.О.",
                  alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "Проверил:",
                  alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "канд. пед. наук, доцент:",
                  alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=FONT_SIZE, space_after=Pt(0))
    add_paragraph(doc, "Фёдорова О.В.",
                  alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=FONT_SIZE, space_after=Pt(0))

    for _ in range(4):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_paragraph(doc, "Казань 2025",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, size=FONT_SIZE)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # СОДЕРЖАНИЕ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Содержание")
    add_toc(doc)
    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ВВЕДЕНИЕ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Введение")

    add_body(doc,
        "В современных условиях информационные технологии являются неотъемлемой частью "
        "ведения бизнеса в сфере предоставления услуг. Фитнес-индустрия не является исключением: "
        "автоматизация процессов управления абонементами, расписанием тренировок, записью клиентов "
        "и мониторингом их прогресса существенно повышает эффективность работы фитнес-центра и "
        "качество обслуживания.")

    add_body(doc,
        "Информационные системы в данной сфере позволяют централизованно управлять данными "
        "о клиентах, абонементах, тренерах и тренировках, а также предоставляют пользователям "
        "удобный инструмент для самостоятельного оформления покупок и записи на занятия через "
        "веб-интерфейс.")

    add_body(doc,
        "Цель данной курсовой работы — разработка полнофункционального веб-приложения "
        "для фитнес-центра. Разрабатываемая система обеспечивает автоматизацию следующих "
        "процессов: управление каталогом абонементов, составление расписания групповых тренировок, "
        "онлайн-запись клиентов на занятия, отслеживание персонального прогресса, а также "
        "администрирование всех перечисленных сущностей.")

    add_body(doc, "Работа включает следующие этапы:")
    add_body(doc, "1. Проектирование архитектуры информационной системы и базы данных.")
    add_body(doc, "2. Разработка серверной части (REST API) и клиентского веб-приложения.")
    add_body(doc, "3. Контейнеризация и развёртывание с использованием Docker.")
    add_body(doc, "4. Тестирование функциональности и документирование результатов.")

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ПОСТАНОВКА ЗАДАЧИ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Постановка задачи")

    add_body(doc,
        "Цель работы: приобретение практических навыков разработки full-stack "
        "веб-приложений с использованием современных технологий серверной и клиентской "
        "разработки, а также контейнеризации.")

    add_body(doc,
        "Для достижения данной цели необходимо создать веб-приложение для фитнес-центра, "
        "которое предоставит пользователям возможность просматривать каталог абонементов, "
        "записываться на групповые тренировки, приобретать абонементы и отслеживать персональный "
        "прогресс. Администраторам система должна обеспечивать полноценное управление контентом.")

    add_body(doc, "Для реализации проекта требуется выполнить следующие задачи:")
    tasks = [
        "Определить программные средства, необходимые для разработки веб-приложения.",
        "Спроектировать архитектуру приложения и базу данных.",
        "Разработать серверную часть (REST API) с аутентификацией, авторизацией и валидацией.",
        "Реализовать клиентскую часть (SPA) с административной панелью и личным кабинетом.",
        "Настроить контейнеризацию с помощью Docker и Docker Compose.",
        "Провести тестирование веб-приложения для выявления и устранения возможных ошибок.",
    ]
    for t in tasks:
        add_body(doc, f"– {t}", indent=True)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ОПИСАНИЕ ПРОГРАММНЫХ СРЕДСТВ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Описание программных средств, используемых при разработке web-приложения")

    add_body(doc,
        "Разработка веб-приложения требует применения различных программных средств, "
        "обеспечивающих его функциональность, безопасность, масштабируемость и удобство "
        "взаимодействия с пользователем.")

    # --- Backend technologies ---
    add_heading3(doc, "3.1. Серверная часть (Backend)")

    backend_techs = [
        ("ASP.NET Core 8.0", "Серверный фреймворк, обеспечивающий высокую производительность при обработке "
         "HTTP-запросов. Обладает встроенной поддержкой внедрения зависимостей (Dependency Injection), "
         "конвейерной обработки запросов (middleware pipeline) и авторизации на основе ролей. "
         "Кроссплатформенность фреймворка позволяет развёртывать приложение как в Windows-среде, "
         "так и в Linux-контейнерах."),
        ("Entity Framework Core 8.0", "ORM-фреймворк с подходом Code-First. Позволяет описывать схему "
         "базы данных непосредственно в коде C# с последующей автоматической генерацией DDL-инструкций. "
         "Fluent API используется для конфигурирования индексов, ограничений уникальности и каскадного удаления."),
        ("PostgreSQL 16", "Реляционная СУБД с открытым исходным кодом. Обладает поддержкой транзакций ACID, "
         "расширенной системой индексов (включая частичные индексы), бесплатным лицензированием и широкой "
         "совместимостью с контейнерными средами."),
        ("JWT (JSON Web Token)", "Механизм аутентификации на основе токенов. Токены подписываются "
         "симметричным ключом по алгоритму HMAC-SHA256 и содержат утверждения (claims) о личности "
         "и роли пользователя, обеспечивая stateless-аутентификацию."),
        ("BCrypt.Net-Next", "Библиотека для хеширования паролей по алгоритму BCrypt. Обеспечивает "
         "устойчивость к атакам перебором за счёт настраиваемого фактора стоимости и автоматической "
         "генерации соли."),
        ("FluentValidation", "Библиотека декларативной валидации входных данных. Правила валидации "
         "описываются в отдельных классах-валидаторах, что обеспечивает разделение ответственности "
         "и тестируемость логики проверки."),
        ("Swashbuckle (Swagger)", "Генератор интерактивной документации API на основе спецификации "
         "OpenAPI. Автоматически формирует описание эндпоинтов, параметров и схем данных."),
    ]
    for name, desc in backend_techs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = LINE_SPACING
        run_name = p.add_run(f"{name} ")
        set_run_font(run_name, bold=True)
        run_desc = p.add_run(f"– {desc}")
        set_run_font(run_desc)

    # --- Frontend technologies ---
    add_heading3(doc, "3.2. Клиентская часть (Frontend)")

    frontend_techs = [
        ("React 18 + TypeScript", "Библиотека для построения пользовательских интерфейсов на основе "
         "компонентного подхода. TypeScript обеспечивает статическую типизацию, обнаружение ошибок "
         "на этапе компиляции и улучшает навигацию по коду."),
        ("Vite 5.0", "Инструмент сборки фронтенд-приложений. Обеспечивает мгновенный запуск "
         "dev-сервера за счёт использования нативных ES-модулей и оптимизированную production-сборку "
         "на основе Rollup."),
        ("Material-UI (MUI) 5.15", "Библиотека готовых UI-компонентов, реализующих принципы Material Design. "
         "Применяются компоненты DataGrid, Dialog, AppBar, Card и другие. Обеспечивает единообразие "
         "интерфейса и сокращение объёма пользовательского CSS."),
        ("React Router 6", "Библиотека клиентской маршрутизации для организации навигации в SPA, "
         "включая вложенные маршруты и защищённые маршруты с проверкой авторизации."),
        ("Axios", "HTTP-клиент для взаимодействия с REST API. Применяются перехватчики (interceptors) "
         "для автоматического добавления JWT-токена и обработки ошибок аутентификации."),
        ("Zustand", "Библиотека управления состоянием с минималистичным API. Применяется для хранения "
         "состояния аутентификации, корзины и уведомлений."),
    ]
    for name, desc in frontend_techs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = LINE_SPACING
        run_name = p.add_run(f"{name} ")
        set_run_font(run_name, bold=True)
        run_desc = p.add_run(f"– {desc}")
        set_run_font(run_desc)

    # --- Infrastructure ---
    add_heading3(doc, "3.3. Инфраструктура")

    infra_techs = [
        ("Docker", "Контейнеризация приложения обеспечивает воспроизводимость среды выполнения. "
         "Используется multi-stage сборка: этап 1 — сборка фронтенда (Node.js), этап 2 — сборка "
         "бэкенда (.NET SDK), этап 3 — финальный образ на базе ASP.NET Runtime."),
        ("Docker Compose", "Оркестрация многоконтейнерного приложения. Описывает два сервиса "
         "(PostgreSQL и приложение) с указанием зависимостей, переменных окружения, health-чеков "
         "и именованных томов для персистентности данных."),
    ]
    for name, desc in infra_techs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = LINE_SPACING
        run_name = p.add_run(f"{name} ")
        set_run_font(run_name, bold=True)
        run_desc = p.add_run(f"– {desc}")
        set_run_font(run_desc)

    add_body(doc,
        "Комплексное использование перечисленных технологий обеспечивает разработку надёжного, "
        "масштабируемого и безопасного веб-приложения, способного эффективно обрабатывать данные "
        "и предоставлять удобный интерфейс пользователям.")

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ЭТАПЫ ПРОЕКТИРОВАНИЯ И РАЗРАБОТКИ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Этапы проектирования и разработки web-приложения")

    add_body(doc,
        "На первом этапе разработки веб-приложения определяются требования к функциональности "
        "системы и проектируются UML-диаграммы. Затем разрабатывается архитектура приложения "
        "и проектируется база данных. На основе этого создаётся программная реализация серверной "
        "и клиентской частей. Завершающим этапом является тестирование.")

    # ── 4.1 Определение требований ──
    add_heading3(doc, "4.1. Определение требований")

    add_body(doc,
        "Основной задачей веб-приложения является предоставление пользователям возможности "
        "просматривать каталог абонементов фитнес-центра, приобретать абонементы, просматривать "
        "расписание групповых тренировок, записываться на занятия и отслеживать персональный прогресс.")

    add_body(doc,
        "Ещё одной важной функцией является администрирование. Приложение должно поддерживать "
        "управление пользователями, абонементами, тренировками, тренерами и категориями "
        "с разграничением прав доступа на основе ролей (RBAC).")

    add_body(doc,
        "На основе требований разработана диаграмма вариантов использования (Use Case), "
        "визуализирующая взаимодействие актёров (Гость, Пользователь, Администратор) с системой. "
        "Диаграмма представлена на рисунке ниже.")

    fig_num = add_figure(
        doc,
        os.path.join(IMAGES_DIR, "usecase.png"),
        "Диаграмма вариантов использования (Use Case)",
        fig_num,
        width=Cm(16)
    )

    add_body(doc,
        "В диаграмме выделены три актёра: Гость (может просматривать каталог и авторизоваться), "
        "Пользователь (наследует возможности гостя; может совершать покупки, записываться на "
        "тренировки и управлять прогрессом) и Администратор (наследует возможности пользователя; "
        "управляет всеми сущностями системы).")

    # ── 4.2 Проектирование архитектуры ──
    add_heading3(doc, "4.2. Проектирование архитектуры")

    add_body(doc,
        "Веб-приложение строится на основе трёхуровневой архитектуры (Three-Tier), "
        "где каждый уровень выполняет определённые функции и обеспечивает структурированное "
        "разделение логики.")

    # List items for architecture layers
    arch_items = [
        ("Слой представления (Presentation Layer)", "отвечает за приём HTTP-запросов "
         "и формирование ответов. Реализован контроллерами ASP.NET Core."),
        ("Слой бизнес-логики (Business Logic Layer)", "обрабатывает запросы, выполняет "
         "валидацию бизнес-правил, координирует операции с данными. Реализован сервисными классами."),
        ("Слой доступа к данным (Data Access Layer)", "управляет хранилищем информации, "
         "выполняя запросы через Entity Framework Core."),
    ]
    for name, desc in arch_items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = LINE_SPACING
        r1 = p.add_run(f"{name} ")
        set_run_font(r1, bold=True)
        r2 = p.add_run(f"– {desc}")
        set_run_font(r2)

    add_body(doc,
        "Каждый уровень взаимодействует с другими через определённые интерфейсы, что позволяет "
        "создавать гибкие и масштабируемые веб-приложения, разделяя функциональность и ответственность "
        "между компонентами.")

    # ── Patterns subsection ──
    add_heading3(doc, "4.2.1. Применённые паттерны проектирования")

    patterns = [
        ("MVC (Model-View-Controller)",
         "Серверная часть реализована на основе паттерна MVC: Model — классы сущностей "
         "и DTO-объекты; Controller — классы, наследующие ControllerBase; View — React-приложение "
         "на стороне клиента."),
        ("DTO (Data Transfer Object)",
         "Для передачи данных между слоями применяются неизменяемые записи (record) языка C#. "
         "DTO изолируют внутреннюю структуру сущностей от API-контракта и предотвращают передачу "
         "конфиденциальных данных клиенту."),
        ("Middleware (Цепочка обязанностей)",
         "Обработка HTTP-запросов организована в виде конвейера middleware-компонентов: "
         "ExceptionMiddleware, UseStaticFiles, Swagger, CORS, Authentication, Authorization, "
         "MapControllers, MapFallbackToFile."),
        ("Dependency Injection (Внедрение зависимостей)",
         "Все сервисы регистрируются в контейнере зависимостей ASP.NET Core и внедряются "
         "в контроллеры через конструктор, обеспечивая слабую связанность компонентов."),
        ("Strategy (Стратегия)",
         "FluentValidation реализует паттерн Strategy: для каждого типа DTO создаётся "
         "отдельный класс-валидатор, фреймворк автоматически подбирает нужный валидатор."),
        ("Observer (Наблюдатель)",
         "Zustand реализует паттерн Observer: компоненты React подписываются на изменения "
         "в хранилище и автоматически перерендериваются при обновлении состояния."),
        ("Interceptor (Перехватчик)",
         "В Axios-клиенте настроены перехватчики запросов (добавление JWT-токена) и ответов "
         "(обработка ошибки 401 с перенаправлением на страницу входа)."),
    ]
    for name, desc in patterns:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = LINE_SPACING
        r1 = p.add_run(f"{name}. ")
        set_run_font(r1, bold=True)
        r2 = p.add_run(desc)
        set_run_font(r2)

    # ── 4.3 Проектирование базы данных ──
    add_heading3(doc, "4.3. Проектирование базы данных")

    add_body(doc,
        "На данном этапе разработана структура базы данных, хранящей информацию о пользователях, "
        "абонементах, тренировках, тренерах, покупках, записях и прогрессе. Для реализации "
        "применяется подход Code-First с использованием Entity Framework Core и Fluent API.")

    add_body(doc,
        "Диаграмма классов (сущностей предметной области) представлена на рисунке ниже.")

    fig_num = add_figure(
        doc,
        os.path.join(IMAGES_DIR, "class-diagram.png"),
        "Диаграмма классов (сущности предметной области)",
        fig_num,
        width=Cm(15)
    )

    # ── 4.4 Разработка приложения ──
    add_heading3(doc, "4.4. Разработка приложения")

    add_body(doc, "Разработка приложения включает следующие направления:")
    dev_items = [
        "Создание моделей данных (User, Membership, Training, Purchase, Booking, Progress и др.).",
        "Разработка REST API с полным набором CRUD-операций для каждой сущности.",
        "Реализация аутентификации и авторизации на основе JWT-токенов и RBAC.",
        "Создание интерфейса пользователя с использованием React и Material-UI.",
        "Разработка административной панели с таблицами (DataGrid) и модальными окнами.",
        "Реализация личного кабинета пользователя с отслеживанием прогресса.",
        "Настройка Docker-контейнеризации для одной командой развёртывания.",
    ]
    for item in dev_items:
        add_body(doc, f"– {item}", indent=True)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ОПИСАНИЕ СТРУКТУРЫ БАЗЫ ДАННЫХ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Описание структуры базы данных")

    add_body(doc,
        "Модель предметной области содержит 10 сущностей и 2 перечисления, "
        "организованных в следующие предметные группы.")

    # --- User ---
    add_heading3(doc, "5.1. Таблица User (Пользователь)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "Guid (PK)", "Уникальный идентификатор"],
            ["Email", "string", "Адрес электронной почты (уникальный индекс)"],
            ["PasswordHash", "string", "Хеш пароля (BCrypt)"],
            ["Role", "string (max 20)", 'Роль: "Admin" или "User"'],
            ["FullName", "string", "Полное имя"],
            ["CreatedAt", "DateTime", "Дата регистрации"],
        ])

    # --- Membership ---
    add_heading3(doc, "5.2. Таблица Membership (Абонемент)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["Name", "string", "Название абонемента"],
            ["Description", "string", "Описание"],
            ["Price", "decimal(10,2)", "Стоимость (индекс для фильтрации)"],
            ["DurationDays", "int", "Срок действия в днях"],
            ["CreatedAt", "DateTime", "Дата создания"],
        ])

    # --- MembershipOption ---
    add_heading3(doc, "5.3. Таблица MembershipOption (Опция абонемента)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["MembershipId", "int (FK)", "Ссылка на абонемент"],
            ["OptionText", "string", "Текст опции"],
        ])

    # --- Category ---
    add_heading3(doc, "5.4. Таблица Category (Категория тренировки)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["Name", "string", "Название категории"],
        ])

    # --- Coach ---
    add_heading3(doc, "5.5. Таблица Coach (Тренер)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["FullName", "string", "ФИО тренера"],
            ["PhotoUrl", "string?", "Путь к фото (загружается из файла)"],
            ["Specialization", "string?", "Специализация"],
        ])

    # --- Training ---
    add_heading3(doc, "5.6. Таблица Training (Групповая тренировка)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["CategoryId", "int (FK)", "Ссылка на категорию (индекс)"],
            ["CoachId", "int (FK)", "Ссылка на тренера"],
            ["Description", "string", "Описание тренировки"],
            ["CoachPhotoUrl", "string?", "Переопределение фото тренера"],
            ["StartTime", "DateTime", "Дата и время начала (индекс)"],
            ["MaxParticipants", "int", "Максимальное число участников"],
        ])

    # --- Purchase ---
    add_heading3(doc, "5.7. Таблица Purchase (Покупка)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["UserId", "Guid (FK)", "Ссылка на пользователя"],
            ["MembershipId", "int (FK)", "Ссылка на абонемент"],
            ["PriceAtPurchase", "decimal(10,2)", "Цена на момент покупки"],
            ["Status", "PurchaseStatus", "Статус: Pending / Paid"],
            ["CreatedAt", "DateTime", "Дата создания"],
        ])

    # --- Booking ---
    add_heading3(doc, "5.8. Таблица Booking (Запись на тренировку)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["TrainingId", "int (FK)", "Ссылка на тренировку"],
            ["UserId", "Guid (FK)", "Ссылка на пользователя"],
            ["Status", "BookingStatus", "Статус: Active / Cancelled"],
        ])

    add_body(doc,
        "Для пары (TrainingId, UserId) задан уникальный фильтрованный индекс с условием "
        "Status = Active, что предотвращает повторную запись на ту же тренировку при наличии "
        "активной записи, но допускает повторную запись после отмены.")

    # --- ProgressTracker ---
    add_heading3(doc, "5.9. Таблица ProgressTracker (Трекер прогресса)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["UserId", "Guid (FK)", "Ссылка на пользователя"],
            ["Title", "string", "Название упражнения/показателя"],
            ["GoalValue", "double", "Целевое значение"],
            ["Unit", "string", "Единица измерения (кг, мин, км)"],
            ["CreatedAt", "DateTime", "Дата создания"],
        ])

    # --- ProgressEntry ---
    add_heading3(doc, "5.10. Таблица ProgressEntry (Запись в трекере)")
    add_table_from_data(doc,
        ["Атрибут", "Тип", "Описание"],
        [
            ["Id", "int (PK)", "Уникальный идентификатор"],
            ["TrackerId", "int (FK)", "Ссылка на трекер"],
            ["Value", "double", "Зафиксированное значение"],
            ["DateRecorded", "DateTime", "Дата замера"],
        ])

    # --- Enumerations ---
    add_heading3(doc, "5.11. Перечисления")
    add_body(doc, "PurchaseStatus: Pending (ожидание оплаты), Paid (оплачено).")
    add_body(doc, "BookingStatus: Active (активная запись), Cancelled (отменена).")

    # --- Relationships ---
    add_heading3(doc, "5.12. Связи между сущностями")

    add_table_from_data(doc,
        ["Связь", "Тип", "Описание"],
        [
            ["User → Purchase", "1 : M", "Пользователь совершает множество покупок"],
            ["User → Booking", "1 : M", "Пользователь записывается на множество тренировок"],
            ["User → ProgressTracker", "1 : M", "Пользователь создаёт множество трекеров"],
            ["Membership → MembershipOption", "1 : M", "Абонемент включает множество опций (каскадное удаление)"],
            ["Membership → Purchase", "1 : M", "Абонемент фигурирует в множестве покупок"],
            ["Category → Training", "1 : M", "Категория объединяет множество тренировок"],
            ["Coach → Training", "1 : M", "Тренер ведёт множество тренировок"],
            ["Training → Booking", "1 : M", "На тренировку записывается множество пользователей"],
            ["ProgressTracker → ProgressEntry", "1 : M", "Трекер содержит множество замеров (каскадное удаление)"],
        ])

    add_body(doc,
        "Связей типа «многие ко многим» (M:M) в явном виде в модели нет. Связь между "
        "пользователем и тренировкой опосредована через сущность Booking, что позволяет "
        "хранить статус записи.")

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ОПИСАНИЕ СТРУКТУРЫ WEB-ПРИЛОЖЕНИЯ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Описание структуры web-приложения")

    add_body(doc,
        "Серверная часть приложения разделена на 7 логических модулей, каждый из которых "
        "включает контроллер (приём HTTP-запросов) и сервис (бизнес-логика).")

    # --- Auth module ---
    add_heading3(doc, "6.1. Модуль аутентификации")

    add_body(doc,
        "Контроллер AuthController (маршрут /api/auth) предоставляет три эндпоинта: "
        "регистрация нового пользователя, авторизация существующего и получение профиля "
        "текущего пользователя по JWT-токену.")

    add_body(doc,
        "Алгоритм регистрации: принимаются email, пароль и ФИО. Выполняется проверка "
        "уникальности email в базе данных. При отсутствии дубля пароль хешируется алгоритмом "
        "BCrypt, создаётся новая запись User с ролью «User». Формируется JWT-токен, содержащий "
        "утверждения (claims): идентификатор пользователя (sub), email, роль и ФИО. Токен "
        "подписывается симметричным ключом HMAC-SHA256 со сроком действия 120 минут.")

    add_body(doc,
        "Алгоритм авторизации: по переданному email осуществляется поиск пользователя "
        "в базе данных. Введённый пароль верифицируется методом BCrypt.Verify. При успешной "
        "верификации генерируется JWT-токен аналогично процедуре регистрации.")

    # --- Memberships module ---
    add_heading3(doc, "6.2. Модуль управления абонементами")

    add_body(doc,
        "Контроллер MembershipsController (маршрут /api/memberships) реализует полный набор "
        "CRUD-операций. Чтение доступно всем пользователям; создание, обновление и удаление — "
        "только администраторам.")

    add_body(doc,
        "Алгоритм получения списка: формируется запрос с включением связанных опций. "
        "Применяются опциональные фильтры: минимальная и максимальная цена, текстовый поиск "
        "по названию. Результат сортируется по цене с серверной пагинацией (Skip/Take). "
        "Возвращается обёртка PagedResult<T>, содержащая элементы текущей страницы и общее "
        "количество записей.")

    add_body(doc,
        "При обновлении абонемента существующие опции полностью удаляются из дочерней таблицы "
        "и пересоздаются на основе переданных данных.")

    # --- Trainings module ---
    add_heading3(doc, "6.3. Модуль управления тренировками, категориями и тренерами")

    add_body(doc,
        "Контроллеры: TrainingsController (/api/trainings), CategoriesController (/api/categories), "
        "CoachesController (/api/coaches). Сервис TrainingService объединяет логику, связанную "
        "с расписанием тренировок, справочником категорий и управлением тренерами.")

    add_body(doc,
        "Алгоритм получения расписания: запрос к таблице Trainings с включением связанных "
        "сущностей Category, Coach и Bookings. Применяются фильтры по категории, тренеру, дате "
        "и текстовому поиску. Для каждой тренировки вычисляется текущее число участников как "
        "количество записей Booking со статусом Active.")

    add_body(doc,
        "Алгоритм загрузки фото тренера: контроллер валидирует файл (допустимые форматы: "
        "jpg, png, webp; максимальный размер: 5 МБ). Сервис определяет путь к каталогу загрузок, "
        "при наличии предыдущего файла удаляет его, сохраняет новый файл с уникальным именем "
        "и обновляет поле PhotoUrl в базе данных.")

    # --- Purchases module ---
    add_heading3(doc, "6.4. Модуль покупок")

    add_body(doc,
        "Контроллер PurchasesController (маршрут /api/purchases) предоставляет эндпоинты "
        "для создания покупки, имитации оплаты, получения списка покупок текущего пользователя "
        "и (для администратора) всех покупок.")

    add_body(doc,
        "Алгоритм создания покупки: по переданному MembershipId выполняется запрос для "
        "получения текущей цены. Создаётся запись Purchase с фиксированной ценой "
        "(PriceAtPurchase) и начальным статусом Pending. Фиксация цены обеспечивает "
        "неизменность суммы при последующем изменении прайса.")

    add_body(doc,
        "Алгоритм имитации оплаты: выполняется поиск покупки по Id с проверкой принадлежности "
        "текущему пользователю. Если покупка уже оплачена, операция отклоняется. "
        "В противном случае статус обновляется на Paid.")

    # --- Bookings module ---
    add_heading3(doc, "6.5. Модуль записи на тренировки")

    add_body(doc,
        "Контроллер BookingsController (маршрут /api/bookings) предоставляет эндпоинты "
        "для записи на тренировку, отмены записи, получения записей пользователя.")

    add_body(doc,
        "Алгоритм создания записи: подсчитывается количество активных записей на тренировку. "
        "Если оно равно или превышает MaxParticipants, операция отклоняется. Далее проверяется "
        "наличие активной записи данного пользователя для предотвращения дублирования. "
        "При успешном прохождении обеих проверок создаётся Booking со статусом Active.")

    add_body(doc,
        "Алгоритм отмены записи: статус обновляется на Cancelled. Фильтрованный уникальный "
        "индекс автоматически исключает отменённые записи, что позволяет пользователю "
        "повторно записаться на ту же тренировку после отмены.")

    # --- Progress module ---
    add_heading3(doc, "6.6. Модуль отслеживания прогресса")

    add_body(doc,
        "Контроллер ProgressController (маршрут /api/progress) предоставляет эндпоинты "
        "для CRUD-операций с трекерами и замерами. Все операции привязаны к текущему пользователю.")

    add_body(doc,
        "Алгоритм получения трекеров с расчётом динамики: загружаются все трекеры пользователя "
        "с включением записей Entries. Для каждого трекера записи сортируются по дате. "
        "Определяется последнее значение (LastValue). Если записей две или более, вычисляется "
        "процент изменения между последним и предпоследним замером по формуле: "
        "changePercent = (last − prev) / prev × 100, результат округляется до одного знака.")

    # --- Admin module ---
    add_heading3(doc, "6.7. Модуль администрирования")

    add_body(doc,
        "Контроллер AdminController (маршрут /api/admin). Все эндпоинты доступны только "
        "пользователям с ролью Admin. Предоставляет получение списка пользователей "
        "с серверной пагинацией и текстовым поиском, а также эндпоинт смены роли "
        "пользователя. Допустимые значения роли ограничены множеством {Admin, User}.")

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ДИАГРАММЫ ПОСЛЕДОВАТЕЛЬНОСТИ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Диаграммы последовательности")

    add_body(doc,
        "В данном разделе представлены диаграммы последовательности для ключевых "
        "API-процессов, в которых логика работы существенно изменяется по сравнению "
        "с типовыми CRUD-операциями.")

    seq_diagrams = [
        ("seq-register.png", "Диаграмма последовательности: регистрация пользователя"),
        ("seq-login.png", "Диаграмма последовательности: авторизация пользователя"),
        ("seq-purchase.png", "Диаграмма последовательности: покупка абонемента"),
        ("seq-booking.png", "Диаграмма последовательности: запись на тренировку"),
        ("seq-coach-photo.png", "Диаграмма последовательности: загрузка фото тренера"),
    ]

    for filename, caption in seq_diagrams:
        fig_num = add_figure(
            doc,
            os.path.join(IMAGES_DIR, filename),
            caption,
            fig_num,
            width=Cm(14)
        )

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # РЕЗУЛЬТАТЫ ТЕСТИРОВАНИЯ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Результаты тестирования с примерами визуальных форм приложения")

    add_body(doc,
        "В данном разделе представлены скриншоты работающего веб-приложения, "
        "демонстрирующие основные пользовательские сценарии.")

    screenshots = [
        "Главная страница",
        "Каталог абонементов",
        "Расписание тренировок",
        "Страница авторизации",
        "Страница регистрации",
        "Личный кабинет пользователя (покупки и записи)",
        "Раздел «Мой прогресс»",
        "Админ-панель: панель управления (Dashboard)",
        "Админ-панель: управление пользователями",
        "Админ-панель: управление абонементами",
        "Админ-панель: управление тренировками",
        "Админ-панель: управление тренерами",
    ]
    for caption in screenshots:
        fig_num = add_screenshot_placeholder(doc, caption, fig_num)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ЗАКЛЮЧЕНИЕ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Заключение")

    add_body(doc,
        "В процессе выполнения курсовой работы было разработано полнофункциональное "
        "веб-приложение для фитнес-центра, реализующее весь заявленный набор функциональных "
        "требований.")

    add_body(doc,
        "Серверная часть приложения реализована на платформе ASP.NET Core 8 с использованием "
        "Entity Framework Core для доступа к данным, FluentValidation для валидации, "
        "JWT-аутентификации для безопасности и ролевой модели RBAC для разграничения доступа. "
        "Клиентская часть построена на React 18 с TypeScript, Material-UI и Zustand.")

    add_body(doc,
        "Были реализованы следующие ключевые подсистемы: каталог абонементов с фильтрацией "
        "и пагинацией, расписание групповых тренировок с записью и проверкой вместимости, "
        "система покупок с имитацией оплаты, модуль отслеживания персонального прогресса "
        "с расчётом динамики, а также административная панель для управления всеми сущностями.")

    add_body(doc,
        "Для развёртывания приложения подготовлена конфигурация Docker Compose, позволяющая "
        "запустить весь стек (PostgreSQL + приложение) одной командой. Multi-stage сборка "
        "Docker-образа объединяет фронтенд и бэкенд в единый контейнер.")

    add_body(doc,
        "Результаты тестирования подтвердили работоспособность всех реализованных функций. "
        "Веб-приложение удовлетворяет поставленным требованиям и может быть использовано "
        "в качестве основы для дальнейшего развития информационной системы фитнес-центра.")

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # СПИСОК ЛИТЕРАТУРЫ
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Список литературы")

    references = [
        "Полуэктова, Н. Р.  Разработка веб-приложений : учебное пособие для вузов / "
        "Н. Р. Полуэктова. — 2-е изд. — Москва : Издательство Юрайт, 2024. — 204 с.",
        "Сысолетин, Е. Г.  Разработка интернет-приложений : учебное пособие для вузов / "
        "Е. Г. Сысолетин, С. Д. Ростунцев. — Москва : Издательство Юрайт, 2022. — 90 с.",
        "Зыков, С. В. Проектирование и разработка корпоративных информационных систем : "
        "учебное пособие / С. В. Зыков. — Москва : Ай Пи Ар Медиа, 2023. — 394 c.",
        "Маркин, А. В.  Программирование на SQL : учебник и практикум для вузов / "
        "А. В. Маркин. — 3-е изд. — Москва : Издательство Юрайт, 2023. — 805 с.",
        "Гладких, Т. В. Разработка функциональных информационных подсистем организации : "
        "учебное пособие / Т. В. Гладких, Е. В. Воронова. — Воронеж, 2023. — 132 с.",
        "Козлова, Е. А. Архитектура и проектирование информационных систем. — Москва, 2021.",
        "Полякова, Л. Н. Основы SQL : учебное пособие / Л. Н. Полякова. — 3-е изд. — "
        "Москва : ИНТУИТ, 2020. — 273 c.",
        "Microsoft. ASP.NET Core documentation. — URL: https://learn.microsoft.com/aspnet/core/",
        "React. Official documentation. — URL: https://react.dev/",
        "PostgreSQL. Official documentation. — URL: https://www.postgresql.org/docs/16/",
        "Docker. Official documentation. — URL: https://docs.docker.com/",
        "Material-UI. Official documentation. — URL: https://mui.com/material-ui/",
    ]
    for i, ref in enumerate(references, 1):
        add_body(doc, f"{i}.\t{ref}", indent=True)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════
    # ПРИЛОЖЕНИЕ 1
    # ════════════════════════════════════════════════════════════════════
    add_heading1(doc, "Приложение 1. Исходный код PlantUML-диаграмм")

    puml_files = [
        ("usecase.puml", "Диаграмма вариантов использования"),
        ("class-diagram.puml", "Диаграмма классов"),
        ("seq-register.puml", "Диаграмма последовательности: регистрация"),
        ("seq-login.puml", "Диаграмма последовательности: авторизация"),
        ("seq-purchase.puml", "Диаграмма последовательности: покупка"),
        ("seq-booking.puml", "Диаграмма последовательности: запись на тренировку"),
        ("seq-coach-photo.puml", "Диаграмма последовательности: загрузка фото тренера"),
    ]

    for filename, title in puml_files:
        filepath = os.path.join(DOCS_DIR, filename)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = LINE_SPACING
        run = p.add_run(f"{filename} — {title}")
        set_run_font(run, bold=True, size=Pt(12))

        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                code = f.read()
            p_code = doc.add_paragraph()
            p_code.paragraph_format.space_after = Pt(2)
            p_code.paragraph_format.line_spacing = 1.0
            run_code = p_code.add_run(code)
            set_run_font(run_code, name="Consolas", size=Pt(9))

    # ── Save ──
    doc.save(OUTPUT_PATH)
    print(f"Документ сохранён: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
