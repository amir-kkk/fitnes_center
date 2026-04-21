import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


def clean_inline(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.strip()


def parse_table_row(line: str) -> list[str]:
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    return [clean_inline(cell) for cell in raw.split("|")]


def is_separator_row(line: str) -> bool:
    candidate = line.strip().replace("|", "").replace("-", "").replace(":", "").strip()
    return candidate == ""


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            doc.add_paragraph("")
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2 and is_separator_row(table_lines[1]):
                headers = parse_table_row(table_lines[0])
                rows = [parse_table_row(r) for r in table_lines[2:]]
                max_cols = max([len(headers)] + [len(r) for r in rows]) if rows else len(headers)

                table = doc.add_table(rows=1, cols=max_cols)
                table.style = "Table Grid"
                for c, value in enumerate(headers):
                    table.rows[0].cells[c].text = value

                for row in rows:
                    tr = table.add_row().cells
                    for c, value in enumerate(row):
                        tr[c].text = value
            else:
                for raw_line in table_lines:
                    doc.add_paragraph(clean_inline(raw_line))
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            level = max(1, min(9, level))
            title = clean_inline(stripped[level:])
            doc.add_heading(title, level=level - 1)
            i += 1
            continue

        if stripped.startswith(">"):
            text = clean_inline(stripped.lstrip(">"))
            p = doc.add_paragraph(text)
            p.style = doc.styles["Intense Quote"] if "Intense Quote" in doc.styles else doc.styles["Normal"]
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            text = clean_inline(re.sub(r"^[-*]\s+", "", stripped))
            doc.add_paragraph(text, style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            text = clean_inline(re.sub(r"^\d+\.\s+", "", stripped))
            doc.add_paragraph(text, style="List Number")
            i += 1
            continue

        if stripped in {"---", "***"}:
            doc.add_paragraph("")
            i += 1
            continue

        doc.add_paragraph(clean_inline(stripped))
        i += 1

    doc.save(docx_path)


if __name__ == "__main__":
    source = Path("e:/Fitness_web/docs/chapter2-design.md")
    target = Path("e:/Fitness_web/docs/chapter2-design.docx")
    convert_markdown_to_docx(source, target)
